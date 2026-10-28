#!/usr/bin/env python3
"""Generate the DA-3 presentation as a PPTX file.

Reads the recommender metrics file and dashboard screenshots, then produces a
12-slide deck in `out/report/`.

Run:
    python3 -m src.utils.gen_slides

Dependencies:
    python-pptx.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

# Allow imports from the project root when running this file directly.
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from config.paths import METRICS_FILE, PROJECT_ROOT as PROJECT, REPORT_DIR

OUT = REPORT_DIR
FILE_BASE = "DA3_RegNo_Ecommerce-Sales-Recommendation"

BLUE = RGBColor(0x1F, 0x3B, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x55, 0x55, 0x55)

TEAM = ("K PRADEEP KUMAR (23BLC1240) · VAHRAN P (23BLC1221) · "
        "ASHVATH KUMAR (23BLC1269) · K.P ABHISHEK (23BLC1292)")


def load_metrics():
    """Read the recommender metrics file and return a flat dict."""
    path = METRICS_FILE
    if not os.path.exists(path):
        return {}
    metrics = {}
    section = "nlp"
    for ln in open(path):
        ln = ln.strip()
        if not ln or ln.startswith("#"):
            continue
        if ln.startswith("[") and ln.endswith("]"):
            section = "baseline"
            continue
        if "=" in ln:
            k, v = ln.split("=", 1)
            metrics[f"{section}_{k}"] = v
    return metrics


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE
    bg.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1))
    tf2 = sub_box.text_frame
    tf2.text = subtitle
    p2 = tf2.paragraphs[0]
    p2.font.size = Pt(18)
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER

    team_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.8))
    tf3 = team_box.text_frame
    tf3.text = TEAM
    p3 = tf3.paragraphs[0]
    p3.font.size = Pt(14)
    p3.font.color.rgb = WHITE
    p3.alignment = PP_ALIGN.CENTER
    return slide


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE
    bg.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(9), Inches(1.5))
    tf = box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    return slide


def add_content_slide(prs, title, bullets, image_path=None, image_left=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(9.2), Inches(0.7))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if image_path and os.path.exists(image_path):
        if image_left:
            left = Inches(0.4)
            text_left = Inches(5.0)
            text_width = Inches(4.6)
        else:
            left = Inches(5.0)
            text_left = Inches(0.4)
            text_width = Inches(4.6)
        try:
            slide.shapes.add_picture(image_path, left, Inches(1.4), width=Inches(4.6))
        except Exception:
            pass
    else:
        text_left = Inches(0.4)
        text_width = Inches(9.2)

    content_box = slide.shapes.add_textbox(text_left, Inches(1.4), text_width, Inches(5.8))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = GRAY
        p.space_after = Pt(10)
    return slide


def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(9.2), Inches(0.7))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    lt = slide.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(4.4), Inches(0.5))
    ltf = lt.text_frame
    ltf.text = left_title
    lp = ltf.paragraphs[0]
    lp.font.size = Pt(22)
    lp.font.bold = True
    lp.font.color.rgb = BLUE

    lb = slide.shapes.add_textbox(Inches(0.4), Inches(1.9), Inches(4.4), Inches(5.3))
    lbf = lb.text_frame
    lbf.word_wrap = True
    for i, item in enumerate(left_items):
        p = lbf.paragraphs[0] if i == 0 else lbf.add_paragraph()
        p.text = item
        p.font.size = Pt(16)
        p.font.color.rgb = GRAY
        p.space_after = Pt(8)

    rt = slide.shapes.add_textbox(Inches(5.2), Inches(1.3), Inches(4.4), Inches(0.5))
    rtf = rt.text_frame
    rtf.text = right_title
    rp = rtf.paragraphs[0]
    rp.font.size = Pt(22)
    rp.font.bold = True
    rp.font.color.rgb = BLUE

    rb = slide.shapes.add_textbox(Inches(5.2), Inches(1.9), Inches(4.4), Inches(5.3))
    rbf = rb.text_frame
    rbf.word_wrap = True
    for i, item in enumerate(right_items):
        p = rbf.paragraphs[0] if i == 0 else rbf.add_paragraph()
        p.text = item
        p.font.size = Pt(16)
        p.font.color.rgb = GRAY
        p.space_after = Pt(8)
    return slide


def main():
    metrics = load_metrics()
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    dashboard_dir = os.path.join(REPORT_DIR, "dashboard")
    nlp_hit = metrics.get("nlp_hit_rate@10", "N/A")
    base_hit = metrics.get("baseline_hit_rate@10", "N/A")
    try:
        improvement = 100 * (float(nlp_hit) - float(base_hit))
    except Exception:
        improvement = 0.0

    # 1. Title
    add_title_slide(prs,
        "E-Commerce Sales Analytics & Product Recommendation",
        "Hadoop · Hive · Spark · MongoDB · Portuguese NLP | BCSE402L DA-3")

    # 2. Problem Statement & Dataset
    add_content_slide(prs, "Problem Statement & Dataset", [
        "Domain: Brazilian Olist e-commerce marketplace",
        "Need: Understand sales trends, geography, payments, delivery and recommendations",
        "Challenge: ~1.07 purchases per customer makes collaborative filtering hard",
        "Dataset: 9 CSVs, ~1.55 M rows, ~122 MB",
        "99,441 orders, 112,650 line items, 93,358 customers",
        "Only 3.00% of customers ever order twice",
    ])

    # 3. Architecture
    add_content_slide(prs, "Pipeline Architecture & Stack", [
        "Stack: Hadoop 3.5.0 · Hive 4.2.1 · Spark 4.2.0 · MongoDB 8.3.11",
        "Raw CSVs → HDFS /olist/raw (immutable)",
        "Clean layer /olist/clean fixes headers, newlines, quotes",
        "MapReduce: revenue by product, orders by month, payment distribution",
        "Hive: external tables + partitioned Parquet warehouse",
        "Spark: curated fact table + cross-validated SQL analytics",
        "Portuguese NLP sentiment enriches the recommender",
        "MongoDB: recommendations + dashboard aggregates",
    ], image_path=os.path.join(REPORT_DIR, "fig1_architecture.png"), image_left=True)

    # 4. MapReduce
    add_two_column_slide(prs, "MapReduce Implementation",
        "Java Jobs",
        [
            "RevenueByProduct: 112,651 → 32,951 records via combiner",
            "OrdersByMonth: cleaning inside mapper, 23 monthly counts",
            "Custom partitioner and Hadoop counters",
        ],
        "Streaming Job",
        [
            "Payment distribution in Python",
            "Separate mapper/combiner/reducer",
            "Outputs count, total and average per payment type",
        ])

    # 5. Hive & Spark
    add_two_column_slide(prs, "Hive Warehouse & Spark ETL",
        "Hive",
        [
            "8 external tables over /olist/clean",
            "Managed Parquet table partitioned by month",
            "10 HQL analytics queries, 0 errors",
            "Window functions and dynamic partitioning",
        ],
        "Spark",
        [
            "Joins 6 tables into 112,650-row fact table",
            "Row counts match sources exactly",
            "6 Spark SQL analyses mirror Hive",
            "Cross-engine validation: numbers agree exactly",
        ])

    # 6. Key Results
    add_content_slide(prs, "Key Business Results", [
        "Peak GMV: R$1,153,364.20 in November 2017",
        "Top category: health_beauty with R$1,412,089.53",
        "Top state: São Paulo (SP) with R$5,769,703.15",
        "Credit card dominates: 76,795 transactions",
        "Repeat customer rate: only 3.00%",
        "On-time deliveries: 4.21★ vs late deliveries: 2.55★",
    ], image_path=os.path.join(dashboard_dir, "overview.png"), image_left=True)

    # 7. NLP Sentiment Pipeline
    add_content_slide(prs, "Portuguese NLP Sentiment Pipeline", [
        "Uses real Portuguese review text (40,950 of 99,224 reviews)",
        "Curated lexicon: 1,260 tokens (720 positive, 540 negative)",
        "Normalisation: lower-case, accent removal, punctuation cleaning",
        "Negation handling: 'nao gostei' flips to negative",
        "Sentiment attributed to products weighted by revenue share per order",
        "Product sentiment scores in [-1, 1] for 18,908 products",
    ])

    # 8. Recommender Comparison
    add_content_slide(prs, "Recommendation System Comparison", [
        "Baseline: popularity + category affinity",
        "NLP model: popularity (60%) + confidence-adjusted sentiment (40%)",
        "Both evaluated with leave-one-out on customers with 2+ items",
        f"Baseline: HitRate@10 = {base_hit} | NLP: HitRate@10 = {nlp_hit}",
        f"NLP improves HitRate@10 by {improvement:.2f} percentage points",
        "Cold-start users fall back to global bestsellers",
    ])

    # 9. Dashboard
    add_content_slide(prs, "Streamlit Dashboard", [
        "Interactive dashboard over MongoDB serving layer",
        "Views: Overview, Sales Trends, Categories, Geography, Delivery, Recommendations",
        "Pre-computed aggregates enable fast load times",
        "Customer recommendation lookup with model metrics",
        "Screenshots included in DA-3 report",
    ], image_path=os.path.join(dashboard_dir, "sales_trends.png"), image_left=False)

    # 10. Evaluation & Limitations
    add_content_slide(prs, "Evaluation & Limitations", [
        "Leave-one-out on 5,334 customers with 2+ delivered items",
        "One held-out set sampled once, used for both models",
        "Popularity, category lists and sentiment computed from training data only",
        "Held-out order reviews excluded from sentiment",
        "After evaluation, deployed MongoDB recommendations are rebuilt from the full dataset",
        "Reviews are order-level → product sentiment is a proxy",
        "Sparse data: ~1.07 purchases per customer",
        "Same-dataset benchmark, not a time-based future estimate",
        "Future work: time-based split, pretrained Portuguese model",
    ])

    # 11. Conclusion
    add_content_slide(prs, "Conclusion", [
        "Complete batch Big Data pipeline + dashboard delivered",
        "Portuguese NLP sentiment pipeline processes real review text",
        "NLP-aware recommender improves over popularity baseline",
        "Pipeline reproducible via run_project.sh",
        "Future: pretrained Portuguese/multilingual sentiment model",
        "Future: Spark Structured Streaming for real-time reviews",
    ])

    # 12. References / Thank You
    add_content_slide(prs, "References & Thank You", [
        "Olist. Brazilian E-Commerce Public Dataset by Olist, v6. Kaggle, 2018.",
        "Apache Hadoop 3.5.0, Hive 4.2.1, Spark 4.2.0 documentation.",
        "MongoDB Community Server 8.3.11 manual.",
        "Streamlit 1.50.0 documentation.",
        "",
        "Questions & Discussion",
        TEAM,
    ])

    out_path = os.path.join(OUT, FILE_BASE + ".pptx")
    prs.save(out_path)
    print("wrote", out_path)
    print("slides:", len(prs.slides))


if __name__ == "__main__":
    main()
