#!/usr/bin/env python3
"""Generate DA-1 reports (DOCX + PDF) for two Big Data Analytics projects."""
import os, textwrap
from PIL import Image as PILImage, ImageDraw, ImageFont

OUT = "/Users/abhishek/Documents/BD PROJECT/DA1"
os.makedirs(OUT, exist_ok=True)

# ---------------- flowchart drawing (Pillow) ----------------
FONT_PATHS = [
    "/System/Library/Fonts/Helvetica.ttc",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/System/Library/Fonts/HelveticaNeue.ttc",
]

def load_font(size):
    for p in FONT_PATHS:
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size)
            except Exception:
                continue
    return ImageFont.load_default()

def draw_flow(steps, path):
    W, box_w, box_h, gap, margin = 1400, 1060, 116, 64, 50
    n = len(steps)
    H = margin * 2 + n * box_h + (n - 1) * gap
    img = PILImage.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    font = load_font(36)
    x0 = (W - box_w) // 2
    x1 = x0 + box_w
    cx = W // 2
    y = margin
    for i, s in enumerate(steps):
        d.rounded_rectangle([x0, y, x1, y + box_h], radius=22,
                            fill="#DCE6F1", outline="#1F3B63", width=4)
        # wrap text to fit
        words, lines, cur = s.split(), [], ""
        for w in words:
            t = (cur + " " + w).strip()
            if d.textlength(t, font=font) <= box_w - 60:
                cur = t
            else:
                lines.append(cur); cur = w
        lines.append(cur)
        th = len(lines) * 44
        ty = y + (box_h - th) // 2
        for ln in lines:
            d.text((cx, ty), ln, font=font, fill="#1F3B63", anchor="mm")
            ty += 44
        if i < n - 1:
            ay = y + box_h
            d.line([cx, ay, cx, ay + gap], fill="#1F3B63", width=6)
            d.polygon([(cx - 16, ay + gap - 22), (cx + 16, ay + gap - 22),
                       (cx, ay + gap)], fill="#1F3B63")
        y += box_h + gap
    img.save(path)

# ---------------- content blocks ----------------
# block types: title, subtitle, meta, h1, h2, p, bullets, numbered, table, image, caption

def project1(flow_png):
    return {
        "file": "DA_RegNo_Twitter-Sentiment-Analysis",
        "blocks": [
            ("title", "Real-Time Sentiment Analytics of Social Media Streams using Spark and MongoDB"),
            ("subtitle", "BCSE402L – Big Data Analytics (TH)  |  DA-1: Topic Selection & Planning"),
            ("meta", "Team Members: [Name 1 – Reg. No., Team Lead], [Name 2 – Reg. No.], [Name 3 – Reg. No.], [Name 4 – Reg. No.]"),
            ("meta", "Submission Date: 4 August 2026"),

            ("h1", "1. Problem Statement"),
            ("p", "Domain: Social Media Analytics (Twitter/X data)."),
            ("p", "Problem (one line): To design a Big Data pipeline that ingests high-velocity tweet streams and classifies their sentiment (positive / negative / neutral) in near real time, enabling live tracking of public opinion."),
            ("p", "Why this problem is meaningful:"),
            ("bullets", [
                "Twitter/X generates ~500 million tweets per day; this volume, velocity and variety exceeds what traditional RDBMS-based batch tools can handle.",
                "Brands, election campaigns and disaster-response agencies need opinion signals within minutes, not after overnight batch jobs.",
                "Real use cases: brand-reputation monitoring, crisis/disaster sentiment tracking, and detecting sudden opinion spikes around events or product launches.",
                "The problem is inherently a Big Data problem (the 3 Vs), making it a natural fit for the tools studied in BCSE402L.",
            ]),

            ("h1", "2. Course Relevance"),
            ("p", "The project maps directly onto four modules of the course syllabus:"),
            ("table",
             ["Course Module", "Concept", "Use in Project"],
             [
                 ["Module 2 – Hadoop & HDFS", "Distributed storage, blocks, replication", "HDFS acts as the data lake for raw tweet archives and the labelled training corpus."],
                 ["Module 4 – NoSQL (MongoDB)", "Document store, JSON documents, CRUD", "MongoDB stores processed tweets and per-minute sentiment aggregates; the schema-less model suits varied tweet fields."],
                 ["Module 5 – Spark", "RDDs, DataFrames, Spark SQL, ETL, MLlib", "Spark Structured Streaming is the processing engine; Spark MLlib trains the sentiment classifier; DataFrames handle ETL."],
                 ["Module 6 – Data Streams", "Stream computing, windowing, sampling, filtering, counting", "Sliding-window counts of positive/negative tweets, spam/retweet filtering, and trending-hashtag detection."],
             ]),

            ("h1", "3. Basic Concepts"),
            ("h2", "3.1 Tools"),
            ("bullets", [
                "Apache Spark Structured Streaming: treats a live stream as an unbounded table; processes data in micro-batches with event-time windows and watermarks.",
                "Apache Kafka / TCP socket: injects the tweet stream at a controlled, repeatable velocity for experiments.",
                "MongoDB: document-oriented NoSQL database; fast writes and flexible schema make it a good serving store for stream results.",
                "HDFS: fault-tolerant distributed storage for the raw dataset and trained model artefacts.",
            ]),
            ("h2", "3.2 Algorithm – Sentiment Classification"),
            ("bullets", [
                "Feature extraction: tokenisation → stop-word removal → TF-IDF / HashingTF vectors (Spark MLlib).",
                "Classifier: Multinomial Naive Bayes trained on 1.6 M labelled tweets — fast, parallelisable, and a strong text baseline; Logistic Regression will be evaluated as an alternative.",
                "Evaluation: accuracy, precision/recall and F1 on a held-out split; stream latency and throughput (tweets/sec) for the pipeline.",
            ]),
            ("h2", "3.3 Workflow Diagram"),
            ("image", flow_png),
            ("caption", "Fig. 1: End-to-end workflow of the real-time sentiment analytics pipeline."),

            ("h1", "4. Approach & Roles"),
            ("h2", "4.1 Methodology"),
            ("numbered", [
                "Data collection – Sentiment140 (1.6 M labelled tweets, Kaggle); replay through a socket/Kafka producer to simulate a live stream (the X API is paywalled).",
                "Environment setup – Apache Spark, MongoDB, Kafka; PySpark client.",
                "Preprocessing – URL/mention normalisation, lower-casing, stop-word removal, tokenisation; cleaned corpus stored on HDFS.",
                "Offline model training – TF-IDF + Naive Bayes PipelineModel in Spark MLlib; persist the trained model.",
                "Streaming pipeline – Structured Streaming job: ingest → clean → score sentiment → per-minute windowed aggregation → write to MongoDB.",
                "Visualisation – Streamlit/Matplotlib dashboard: live sentiment gauge, positive/negative trend lines, top hashtags.",
                "Evaluation & documentation – accuracy/latency measurements; DA-2 and DA-3 deliverables.",
            ]),
            ("h2", "4.2 Tentative Timeline"),
            ("table",
             ["Phase", "Work", "Duration"],
             [
                 ["Phase 1", "Environment setup, dataset download, exploration", "4 – 10 Aug"],
                 ["Phase 2", "Preprocessing and HDFS ingestion", "11 – 24 Aug"],
                 ["Phase 3", "Offline model training and evaluation", "25 Aug – 7 Sep"],
                 ["Phase 4", "Streaming pipeline + MongoDB sink (DA-2: 19 Sep)", "8 – 19 Sep"],
                 ["Phase 5", "Dashboard, tuning, documentation", "20 Sep – 12 Oct"],
                 ["Phase 6", "Final report + presentation (DA-3: 20 Oct)", "13 – 20 Oct"],
             ]),
            ("h2", "4.3 Team Roles"),
            ("table",
             ["Member", "Responsibility"],
             [
                 ["Member 1 – [Name] (Lead)", "Overall architecture, Spark streaming pipeline, integration"],
                 ["Member 2 – [Name]", "Data preprocessing, HDFS ingestion, dataset documentation"],
                 ["Member 3 – [Name]", "MLlib model training and evaluation"],
                 ["Member 4 – [Name]", "MongoDB serving layer, dashboard, report and presentation"],
             ]),

            ("h1", "5. Feasibility"),
            ("h2", "5.1 Dataset Sources"),
            ("bullets", [
                "Sentiment140 (Kaggle): 1.6 M tweets with polarity labels, CSV ≈ 240 MB – identified and downloaded.",
                "Optional live source: X/Twitter filtered-stream API (fallback to replay because of the paywall).",
            ]),
            ("h2", "5.2 Evidence of Setup (Initial Progress)"),
            ("bullets", [
                "Apache Spark 3.x installed and verified (PySpark word-count smoke test).",
                "MongoDB Community Server running locally; Spark–MongoDB connector tested.",
                "Dataset downloaded and schema inspected (6 columns: target, id, date, query, user, text).",
                "Git repository initialised for team collaboration.",
            ]),
            ("h2", "5.3 Expected Challenges & Mitigation"),
            ("table",
             ["Challenge", "Mitigation"],
             [
                 ["X API is paid", "Replay Sentiment140 via a socket/Kafka producer at configurable speed"],
                 ["Laptop resource limits", "local[2] mode, micro-batch trigger tuning, sampling"],
                 ["Sarcasm / noisy text lowers accuracy", "Strong preprocessing; compare with Logistic Regression; report baseline honestly"],
                 ["Connector / version issues", "Pin versions; Docker Compose for Kafka + MongoDB"],
                 ["Data skew (retweet storms)", "Filter retweets; watermarking; windowed de-duplication"],
             ]),
            ("p", "The dataset is public, all tools are open-source, and a working local setup already exists — the project is feasible within the semester timeline."),
        ],
    }

def project2(flow_png):
    return {
        "file": "DA_RegNo_Ecommerce-Sales-Recommendation",
        "blocks": [
            ("title", "E-Commerce Sales Analytics and Product Recommendation using Hadoop, Hive and Spark"),
            ("subtitle", "BCSE402L – Big Data Analytics (TH)  |  DA-1: Topic Selection & Planning"),
            ("meta", "Team Members: [Name 1 – Reg. No., Team Lead], [Name 2 – Reg. No.], [Name 3 – Reg. No.], [Name 4 – Reg. No.]"),
            ("meta", "Submission Date: 4 August 2026"),

            ("h1", "1. Problem Statement"),
            ("p", "Domain: E-Commerce Analytics (Brazilian marketplace transaction data – Olist)."),
            ("p", "Problem (one line): To build a batch Big Data analytics pipeline over multi-CSV e-commerce transaction data that uncovers sales trends and customer behaviour and generates personalised product recommendations at a scale where single-machine tools fail."),
            ("p", "Why this problem is meaningful:"),
            ("bullets", [
                "Recommendations drive a large share of e-commerce revenue (industry estimates go up to 35%); small marketplaces lack affordable analytics stacks.",
                "Transaction data arrives as multiple large relational CSVs (orders, items, payments, reviews) whose joins and aggregations outgrow in-memory tools as volume grows.",
                "Insights such as top categories, delivery delays, payment preferences and demand seasonality directly support inventory and marketing decisions.",
                "It is a classic volume + variety Big Data problem, ideal for the Hadoop ecosystem tools taught in BCSE402L.",
            ]),

            ("h1", "2. Course Relevance"),
            ("p", "The project maps directly onto four modules of the course syllabus:"),
            ("table",
             ["Course Module", "Concept", "Use in Project"],
             [
                 ["Module 2 – Hadoop, HDFS, Hive, Pig", "HDFS storage, HQL, partitioning, Pig Latin", "HDFS data lake for all CSVs; Hive external tables partitioned by purchase month; Pig scripts for cleaning."],
                 ["Module 3 – MapReduce", "Map/Reduce, key-value pairs, combiners, partitioners", "Custom MR jobs: revenue per product category, orders per state, payment-type distribution (combiners for local aggregation)."],
                 ["Module 4 – NoSQL (MongoDB)", "Document store", "Serving layer storing per-user recommendation lists for fast lookup."],
                 ["Module 5 – Spark", "DataFrames, Spark SQL, MLlib, ETL", "Multi-table joins and feature preparation; ALS collaborative-filtering recommender; reporting DataFrames."],
             ]),

            ("h1", "3. Basic Concepts"),
            ("h2", "3.1 Tools"),
            ("bullets", [
                "HDFS: splits files into 128 MB replicated blocks; gives horizontal scalability and fault tolerance.",
                "MapReduce: mappers emit (key, value) pairs, e.g. (category, revenue); shuffle groups by key; reducers aggregate; combiners cut network traffic.",
                "Apache Hive: SQL-like warehouse layer (HQL) over HDFS with metastore, partitioning and bucketing.",
                "Apache Spark + MLlib: in-memory DataFrames for joins/ETL; ALS matrix factorisation for recommendations.",
                "MongoDB: stores final recommendation documents of the form {user_id, [product_ids]}.",
            ]),
            ("h2", "3.2 Algorithm – ALS Collaborative Filtering"),
            ("bullets", [
                "A ratings matrix is built from purchase counts and review scores (users × products).",
                "ALS alternates between fixing user and item factor matrices to minimise regularised least-squares error — naturally parallel and supported by Spark MLlib.",
                "Evaluation: RMSE on held-out ratings and top-N hit-rate; popularity-based fallback for cold-start users.",
            ]),
            ("h2", "3.3 Workflow Diagram"),
            ("image", flow_png),
            ("caption", "Fig. 1: End-to-end batch analytics and recommendation workflow."),

            ("h1", "4. Approach & Roles"),
            ("h2", "4.1 Methodology"),
            ("numbered", [
                "Data collection – Brazilian E-Commerce Public Dataset (Olist, Kaggle): 9 CSVs, ~100 k orders (2016–2018), ≈ 120 MB; optional synthetic scale-up to the GB range.",
                "Environment setup – Hadoop pseudo-distributed (HDFS + YARN), Hive, Spark, MongoDB.",
                "Ingestion & cleaning – hdfs dfs -put; Pig/MR cleaning: null handling, type casting, de-duplication, timestamp parsing.",
                "MapReduce analytics – revenue by category, orders by state, payment distribution, monthly order counts.",
                "Hive warehouse – external/partitioned tables; HQL queries for monthly GMV trend, top categories, delivery-delay analysis, repeat-customer rate.",
                "Spark ETL + recommendation – join orders/items/payments/reviews; build ratings; train ALS; generate top-10 recommendations per user.",
                "Serving & visualisation – recommendations written to MongoDB; Matplotlib/Streamlit dashboard for trends and sample recommendations.",
                "Evaluation & documentation – RMSE/hit-rate; DA-2 and DA-3 deliverables.",
            ]),
            ("h2", "4.2 Tentative Timeline"),
            ("table",
             ["Phase", "Work", "Duration"],
             [
                 ["Phase 1", "Setup (Hadoop/Hive/Spark/MongoDB), dataset download", "4 – 10 Aug"],
                 ["Phase 2", "HDFS ingestion + Pig/MR cleaning", "11 – 24 Aug"],
                 ["Phase 3", "MapReduce analytics jobs + Hive warehouse", "25 Aug – 7 Sep"],
                 ["Phase 4", "Spark ETL + ALS recommender (DA-2: 19 Sep)", "8 – 19 Sep"],
                 ["Phase 5", "MongoDB serving + dashboard + tuning", "20 Sep – 12 Oct"],
                 ["Phase 6", "Final report + presentation (DA-3: 20 Oct)", "13 – 20 Oct"],
             ]),
            ("h2", "4.3 Team Roles"),
            ("table",
             ["Member", "Responsibility"],
             [
                 ["Member 1 – [Name] (Lead)", "Hadoop/Hive cluster setup, overall architecture, integration"],
                 ["Member 2 – [Name]", "Data ingestion and cleaning (Pig / MapReduce)"],
                 ["Member 3 – [Name]", "MapReduce + HQL analytics, results documentation"],
                 ["Member 4 – [Name]", "Spark ALS recommender, MongoDB serving layer, dashboard"],
             ]),

            ("h1", "5. Feasibility"),
            ("h2", "5.1 Dataset Sources"),
            ("bullets", [
                "Brazilian E-Commerce Public Dataset by Olist (Kaggle, public licence): 9 relational CSVs, ~100 k orders, ≈ 120 MB – identified and downloaded.",
                "Growth path if more volume is needed: synthetic order generator or a larger public clickstream dataset (e.g. REES46).",
            ]),
            ("h2", "5.2 Evidence of Setup (Initial Progress)"),
            ("bullets", [
                "Hadoop single-node (pseudo-distributed) cluster running; hdfs dfs -ls verified.",
                "Hive metastore initialised; sample HQL queries executed on the dataset.",
                "Spark + PySpark verified; MongoDB Community Server running.",
                "Dataset CSVs inspected (schemas, row counts); Git repository initialised.",
            ]),
            ("h2", "5.3 Expected Challenges & Mitigation"),
            ("table",
             ["Challenge", "Mitigation"],
             [
                 ["Single-node hardware limits", "Pseudo-distributed config; develop on samples, full run for final results"],
                 ["Skewed keys (a few big categories)", "Combiners and a custom partitioner"],
                 ["Complex joins in raw MapReduce", "Do joins in Hive/Spark; keep MR for aggregations only"],
                 ["Cold-start users in ALS", "Popularity-based fallback recommendations"],
                 ["Hadoop/Hive/Spark version compatibility", "Pin compatible versions; document the setup steps"],
             ]),
            ("p", "The dataset is public, the full stack is open-source, and the cluster is already running on a single machine — the project is feasible within the semester timeline."),
        ],
    }

# ---------------- DOCX renderer ----------------
def render_docx(blocks, path):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for s in ("Normal",):
        st = doc.styles[s]; st.font.name = "Calibri"; st.font.size = Pt(11)
    BLUE = RGBColor(0x1F, 0x3B, 0x63)

    def heading(text, size, before=10, after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = BLUE
        return p

    for b in blocks:
        kind = b[0]
        if kind == "title":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b[1]); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = BLUE
        elif kind == "subtitle":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b[1]); r.bold = True; r.font.size = Pt(11)
        elif kind == "meta":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b[1]); r.font.size = Pt(10); r.italic = True
        elif kind == "h1":
            heading(b[1], 13)
        elif kind == "h2":
            heading(b[1], 11.5, before=6, after=2)
        elif kind == "p":
            doc.add_paragraph(b[1])
        elif kind == "bullets":
            for it in b[1]:
                doc.add_paragraph(it, style="List Bullet")
        elif kind == "numbered":
            for it in b[1]:
                doc.add_paragraph(it, style="List Number")
        elif kind == "table":
            headers, rows = b[1], b[2]
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Table Grid"
            for j, h in enumerate(headers):
                c = t.rows[0].cells[j]; c.text = ""
                r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10)
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    c = t.rows[i + 1].cells[j]; c.text = ""
                    r = c.paragraphs[0].add_run(val); r.font.size = Pt(10)
            doc.add_paragraph()
        elif kind == "image":
            doc.add_picture(b[1], width=Inches(5.6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "caption":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b[1]); r.italic = True; r.font.size = Pt(9)
    doc.save(path)

# ---------------- PDF renderer ----------------
def render_pdf(blocks, path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    Image, Table, TableStyle)
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader

    BLUE = HexColor("#1F3B63")
    st_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15,
                              leading=19, alignment=1, textColor=BLUE, spaceAfter=4)
    st_sub = ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=10,
                            leading=13, alignment=1, spaceAfter=2)
    st_meta = ParagraphStyle("m", fontName="Helvetica-Oblique", fontSize=8.5,
                             leading=11, alignment=1, spaceAfter=2)
    st_h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=12.5,
                           leading=15, textColor=BLUE, spaceBefore=9, spaceAfter=4)
    st_h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.5,
                           leading=13, textColor=BLUE, spaceBefore=5, spaceAfter=2)
    st_body = ParagraphStyle("b", fontName="Helvetica", fontSize=9.5,
                             leading=12.5, spaceAfter=4)
    st_bul = ParagraphStyle("bl", parent=st_body, leftIndent=14, spaceAfter=2)
    st_cap = ParagraphStyle("c", fontName="Helvetica-Oblique", fontSize=8.5,
                            leading=11, alignment=1, spaceAfter=6)

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18*mm,
                            rightMargin=18*mm, topMargin=16*mm, bottomMargin=16*mm)
    avail_w = A4[0] - 36*mm
    story = []
    for b in blocks:
        kind = b[0]
        if kind == "title":
            story.append(Paragraph(b[1], st_title))
        elif kind == "subtitle":
            story.append(Paragraph(b[1], st_sub))
        elif kind == "meta":
            story.append(Paragraph(b[1], st_meta))
        elif kind == "h1":
            story.append(Paragraph(b[1], st_h1))
        elif kind == "h2":
            story.append(Paragraph(b[1], st_h2))
        elif kind == "p":
            story.append(Paragraph(b[1], st_body))
        elif kind == "bullets":
            for it in b[1]:
                story.append(Paragraph("• " + it, st_bul))
        elif kind == "numbered":
            for i, it in enumerate(b[1], 1):
                story.append(Paragraph(f"{i}. {it}", st_bul))
        elif kind == "table":
            headers, rows = b[1], b[2]
            ncol = len(headers)
            if ncol == 3 and headers[0] == "Phase":
                widths = [avail_w*0.14, avail_w*0.66, avail_w*0.20]
            elif ncol == 3:
                widths = [avail_w*0.26, avail_w*0.28, avail_w*0.46]
            else:
                widths = [avail_w*0.35, avail_w*0.65]
            cell = ParagraphStyle("tc", fontName="Helvetica", fontSize=8.5, leading=11)
            cellb = ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=8.5, leading=11)
            data = [[Paragraph(h, cellb) for h in headers]] + \
                   [[Paragraph(v, cell) for v in row] for row in rows]
            t = Table(data, colWidths=widths)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#7F7F7F")),
                ("BACKGROUND", (0, 0), (-1, 0), HexColor("#DCE6F1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t); story.append(Spacer(1, 6))
        elif kind == "image":
            iw, ih = ImageReader(b[1]).getSize()
            w = avail_w * 0.62
            story.append(Image(b[1], width=w, height=w * ih / iw, hAlign="CENTER"))
        elif kind == "caption":
            story.append(Paragraph(b[1], st_cap))
    doc.build(story)

# ---------------- main ----------------
if __name__ == "__main__":
    f1 = os.path.join(OUT, "flow_p1.png")
    f2 = os.path.join(OUT, "flow_p2.png")
    draw_flow([
        "Sentiment140 Dataset (Kaggle, 1.6 M tweets)",
        "HDFS Data Lake + Offline Model Training (TF-IDF + Naive Bayes)",
        "Stream Replay via Kafka / TCP Socket",
        "Spark Structured Streaming: Cleaning + Tokenisation",
        "Real-Time Sentiment Scoring (MLlib Model)",
        "Windowed Aggregations + Trending Hashtag Detection",
        "MongoDB Serving Store",
        "Live Dashboard (Trends, Gauges, Alerts)",
    ], f1)
    draw_flow([
        "Olist E-Commerce CSVs (Kaggle, ~100 k orders)",
        "HDFS Data Lake (hdfs dfs -put)",
        "Pig / MapReduce Cleaning + Aggregation Jobs",
        "Hive Warehouse (Partitioned Tables, HQL Analytics)",
        "Spark SQL Joins + Feature Preparation",
        "ALS Recommendation Model (Spark MLlib)",
        "MongoDB Serving Store (Top-10 per User)",
        "Dashboard + Reports (Trends, Sample Recommendations)",
    ], f2)

    for proj in (project1(f1), project2(f2)):
        docx_path = os.path.join(OUT, proj["file"] + ".docx")
        pdf_path = os.path.join(OUT, proj["file"] + ".pdf")
        render_docx(proj["blocks"], docx_path)
        render_pdf(proj["blocks"], pdf_path)
        print("wrote", docx_path)
        print("wrote", pdf_path)
